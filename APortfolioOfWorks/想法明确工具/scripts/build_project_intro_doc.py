from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT_DIR = Path(__file__).resolve().parents[1] / "outputs"
DOCX_PATH = OUT_DIR / "IdeaPilot作品介绍文档.docx"
MD_PATH = OUT_DIR / "IdeaPilot作品介绍文档.md"

INK = RGBColor(17, 24, 39)
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(75, 85, 99)
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
BORDER = "D1D5DB"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders(table, color=BORDER):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_run_font(run, size=None, color=None, bold=None, name="Arial Unicode MS"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold


def add_para(doc, text="", style=None, size=11, color=INK, bold=False, after=6, before=0, align=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    if align is not None:
        p.alignment = align
    if text:
        run = p.add_run(text)
        set_run_font(run, size=size, color=color, bold=bold)
    return p


def add_heading(doc, text, level=1):
    style = f"Heading {level}"
    p = add_para(doc, text, style=style, size={1: 16, 2: 13, 3: 12}[level],
                 color=BLUE if level < 3 else DARK_BLUE, bold=True,
                 before={1: 16, 2: 12, 3: 8}[level], after={1: 8, 2: 6, 3: 4}[level])
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    run = p.add_run(text)
    set_run_font(run, size=11, color=INK)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    run = p.add_run(text)
    set_run_font(run, size=11, color=INK)
    return p


def add_callout(doc, label, text):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    set_table_borders(table, "CBD5E1")
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F4F6F9")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(label)
    set_run_font(r, size=10.5, color=DARK_BLUE, bold=True)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.10
    r2 = p2.add_run(text)
    set_run_font(r2, size=11, color=INK)
    add_para(doc, "", after=4)


def configure_styles(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial Unicode MS"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for level, size, color, before, after in [
        (1, 16, BLUE, 16, 8),
        (2, 13, BLUE, 12, 6),
        (3, 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[f"Heading {level}"]
        style.font.name = "Arial Unicode MS"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_cover(doc):
    add_para(doc, "作品介绍文档", size=11, color=MUTED, bold=True, after=18)
    title = add_para(doc, "IdeaPilot 想法明确工具", size=24, color=INK, bold=True, after=6)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_para(doc, "把模糊项目想法变成可执行 MVP 的 AI 引导工具", size=14, color=MUTED, after=18)

    meta = doc.add_table(rows=4, cols=2)
    set_table_geometry(meta, [1500, 7860])
    set_table_borders(meta, "E5E7EB")
    rows = [
        ("姓名", "王惠诚"),
        ("日期", "2026年7月17日"),
        ("项目状态", "已完成可运行 Demo，可现场演示核心流程"),
        ("一句话", "它像一个想法导航员，帮助用户从“我有个想法”走到“我知道先做什么”。"),
    ]
    for i, (label, value) in enumerate(rows):
        for cell in meta.rows[i].cells:
            set_cell_shading(cell, LIGHT_GRAY if i == 0 else "FFFFFF")
        p1 = meta.cell(i, 0).paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        set_run_font(p1.add_run(label), size=10.5, color=MUTED, bold=True)
        p2 = meta.cell(i, 1).paragraphs[0]
        p2.paragraph_format.space_after = Pt(0)
        set_run_font(p2.add_run(value), size=10.5, color=INK)

    add_callout(
        doc,
        "介绍摘要",
        "IdeaPilot 不是通用聊天机器人，也不是项目管理软件。它专门解决项目开始前最容易混乱的一步：让用户按顺序说清楚目标用户、使用场景、真实需求、MVP 范围和开发提示词。",
    )


def add_overview(doc):
    add_heading(doc, "1. 我的作品是什么", 1)
    add_para(
        doc,
        "我的作品是一个帮助用户把模糊项目想法整理清楚，并生成 AI 开发提示词的工具。它的名字叫 IdeaPilot，也可以叫“想法明确工具”。",
    )
    add_para(
        doc,
        "它面向有项目想法、懂一点开发或愿意使用 AI 编程工具，但不知道怎样把想法说清楚的人。用户在准备做 Demo、课程作业或个人项目之前，可以打开它，用一组分阶段问题把项目范围收敛下来。",
    )

    add_heading(doc, "2. 我想解决的问题", 1)
    add_bullet(doc, "很多人一开始想法很多，但说不清第一版最重要的目标。")
    add_bullet(doc, "有人还没有判断真实需求，就直接让 AI 写代码或做页面。")
    add_bullet(doc, "功能越做越多，范围越来越乱，最后 Demo 不像一个清楚的产品。")
    add_para(
        doc,
        "所以我先做的不是“更强的开发工具”，而是一个帮人明确想法的工具。它解决的是开发前最容易被忽略的一步：先想清楚，再开始做。",
    )

    add_heading(doc, "3. 第一版我先做了什么", 1)
    add_para(doc, "第一版只抓住一个核心动作：把问题问清楚。用户会按五个阶段完成一次项目澄清。")
    stages = [
        ("项目初诊", "先判断想法、目标用户、使用场景和核心问题是否清楚。"),
        ("真实需求判断", "追问问题频率、现有做法、替代方案不足和外部证据。"),
        ("伪需求审查", "检查项目是否只是自己觉得有用，或是否可以被通用 AI、表格、备忘录等工具轻易替代。"),
        ("MVP 收敛", "确定第一版必须做什么，并明确砍掉账号、模板市场、多人协作等非核心功能。"),
        ("提示词生成", "输出一份可直接交给 Codex、Claude Code 等 AI 编程工具继续开发的提示词。"),
    ]
    table = doc.add_table(rows=1 + len(stages), cols=2)
    set_table_geometry(table, [2200, 7160])
    set_table_borders(table)
    headers = ("阶段", "作用")
    for i, header in enumerate(headers):
        cell = table.cell(0, i)
        set_cell_shading(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(header), size=10.5, color=DARK_BLUE, bold=True)
    for row_index, (name, desc) in enumerate(stages, start=1):
        p0 = table.cell(row_index, 0).paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        set_run_font(p0.add_run(name), size=10.5, color=INK, bold=True)
        p1 = table.cell(row_index, 1).paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        set_run_font(p1.add_run(desc), size=10.5, color=INK)

    add_heading(doc, "4. AI 在作品里的作用", 1)
    add_para(doc, "我让 AI 帮我承担“产品教练”的角色。第一版它主要做三件事：")
    add_bullet(doc, "提炼信息：从用户回答里抽取目标用户、场景、问题和限制。")
    add_bullet(doc, "判断风险：提醒证据不足、范围过大、替代方案太强等问题。")
    add_bullet(doc, "推动下一步：根据当前薄弱点提出下一题，让用户继续补齐事实。")
    add_para(doc, "后来我让它改成更严格的诊断流程，而不是普通聊天。这样它不会只顺着用户说，而是会追问、判断，并帮助收敛。")


def add_demo_and_future(doc):
    add_heading(doc, "5. 我会怎样演示 Demo", 1)
    add_para(doc, "现场演示时，我会打开本地运行的 IdeaPilot Demo，然后按下面的路径展示：")
    for step in [
        "输入一个模糊项目想法。",
        "回答目标用户、使用场景和核心问题。",
        "查看 AI 给出的诊断摘要和风险判断。",
        "继续回答真实需求、伪需求审查和 MVP 收敛问题。",
        "最后看到项目定位、MVP 范围、砍掉功能列表和可开发提示词。",
    ]:
        add_number(doc, step)
    add_para(doc, "大家会看到的重点不是页面点击本身，而是它怎样一步步帮我把想法变清楚。")

    add_heading(doc, "6. 当前已经完成的内容", 1)
    add_bullet(doc, "本地可运行的单页 Web 工具。")
    add_bullet(doc, "左侧显示五个项目澄清阶段。")
    add_bullet(doc, "中间逐题提问，每次只让用户回答一个核心问题。")
    add_bullet(doc, "右侧显示当前诊断、风险等级、缺失信息和下一步问题。")
    add_bullet(doc, "完成后可以复制或下载最终的项目澄清报告。")

    add_heading(doc, "7. 下一版我最想加入什么", 1)
    add_para(doc, "如果继续做下一版，我最想加入“保存多个项目”和“导出文档”的能力。")
    add_bullet(doc, "保存多个项目：用户可以回看、修改和继续推进不同想法。")
    add_bullet(doc, "导出文档：把澄清结果导出成 Word、PPT 或可提交的项目说明。")
    add_bullet(doc, "拆分开发任务：把最终提示词进一步拆成页面、接口和测试任务。")
    add_para(doc, "因为这样可以让一次性的想法诊断，升级成持续的项目规划助手，让用户的项目更容易从想法走向完成。")


def add_fill_in_answers(doc):
    doc.add_section(WD_SECTION.NEW_PAGE)
    add_heading(doc, "按图片问题填写的版本", 1)
    add_para(doc, "下面这部分可以直接作为“我的作品介绍稿”的答案使用。")

    answers = [
        ("我可以这样介绍我的作品：", "我的作品叫 IdeaPilot，是一个帮助人把模糊项目想法整理成可执行 MVP 的 AI 引导工具。"),
        ("我的作品是一个 ____。", "想法明确工具 / AI 项目澄清工具。"),
        ("我想帮 ____，解决 ____ 的问题。", "准备做项目但需求还不清楚的人；第一版到底该做什么、哪些功能应该先砍掉。"),
        ("这次我先做了 ____。", "一个能本地运行的单页 Demo，包含分阶段提问、AI 诊断、MVP 收敛和提示词输出。"),
        ("我让 AI 帮我 ____，第一版它 ____，", "判断项目想法是否清楚；会根据回答继续追问，并指出风险。"),
        ("后来我让它改成 ____。", "更严格的产品教练流程，不只是聊天，而是会判断伪需求、要求补充证据并收敛 MVP。"),
        ("现在我来演示 Demo，我会点开 ____，", "想法明确工具的本地网页。"),
        ("大家会看到 ____。", "我输入一个模糊想法后，系统一步步追问，最后生成项目定位、MVP 范围和可开发提示词。"),
        ("如果继续做下一版，我最想加入 ____，", "保存多个项目、导出文档和拆分开发任务。"),
        ("因为这样可以让 ____ 更 ____。", "一个项目想法；清楚、可推进、可继续开发。"),
    ]
    table = doc.add_table(rows=1 + len(answers), cols=2)
    set_table_geometry(table, [3600, 5760])
    set_table_borders(table)
    for idx, header in enumerate(("问题", "可填写答案")):
        cell = table.cell(0, idx)
        set_cell_shading(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(header), size=10.5, color=DARK_BLUE, bold=True)
    for row_index, (question, answer) in enumerate(answers, start=1):
        q = table.cell(row_index, 0).paragraphs[0]
        q.paragraph_format.space_after = Pt(0)
        set_run_font(q.add_run(question), size=10.5, color=INK, bold=True)
        a = table.cell(row_index, 1).paragraphs[0]
        a.paragraph_format.space_after = Pt(0)
        set_run_font(a.add_run(answer), size=10.5, color=INK)

    add_heading(doc, "一段完整口播稿", 1)
    add_para(
        doc,
        "大家好，我的作品叫 IdeaPilot，也就是想法明确工具。它是一个帮助用户把模糊项目想法整理清楚，并生成 AI 开发提示词的工具。我想帮助准备做项目但需求还不清楚的人，解决“第一版到底该做什么、哪些功能应该先砍掉”的问题。",
    )
    add_para(
        doc,
        "这次我先做了一个能本地运行的单页 Demo。它会按项目初诊、真实需求判断、伪需求审查、MVP 收敛和提示词生成这五个阶段一步步提问。我让 AI 帮我做产品教练，第一版它会根据用户回答继续追问，并指出目标用户不清楚、证据不足、范围过大等风险。后来我又让它改成更严格的诊断流程，不只是陪聊，而是会判断和收敛。",
    )
    add_para(
        doc,
        "现在我来演示 Demo。我会点开本地网页，输入一个模糊的项目想法。大家会看到系统怎样一步步追问，最后输出项目定位、MVP 范围、砍掉功能列表和可直接交给 AI 编程工具的开发提示词。如果继续做下一版，我最想加入保存多个项目、导出文档和拆分开发任务，因为这样可以让一个项目想法更清楚、更可推进，也更容易继续开发。",
    )


def build_markdown():
    return """# IdeaPilot 想法明确工具作品介绍

姓名：王惠诚
日期：2026年7月17日

## 一句话介绍

IdeaPilot 是一个帮助用户把模糊项目想法整理清楚，并生成 AI 开发提示词的工具。它像一个想法导航员，帮助用户从“我有个想法”走到“我知道先做什么”。

## 按图片问题填写

1. 我可以这样介绍我的作品：我的作品叫 IdeaPilot，是一个帮助人把模糊项目想法整理成可执行 MVP 的 AI 引导工具。
2. 我的作品是一个想法明确工具 / AI 项目澄清工具。
3. 我想帮准备做项目但需求还不清楚的人，解决第一版到底该做什么、哪些功能应该先砍掉的问题。
4. 这次我先做了一个能本地运行的单页 Demo，包含分阶段提问、AI 诊断、MVP 收敛和提示词输出。
5. 我让 AI 帮我判断项目想法是否清楚，第一版它会根据回答继续追问，并指出风险。
6. 后来我让它改成更严格的产品教练流程，不只是聊天，而是会判断伪需求、要求补充证据并收敛 MVP。
7. 现在我来演示 Demo，我会点开想法明确工具的本地网页。
8. 大家会看到我输入一个模糊想法后，系统一步步追问，最后生成项目定位、MVP 范围和可开发提示词。
9. 如果继续做下一版，我最想加入保存多个项目、导出文档和拆分开发任务。
10. 因为这样可以让一个项目想法更清楚、更可推进、可继续开发。
"""


def main():
    OUT_DIR.mkdir(exist_ok=True)
    doc = Document()
    configure_styles(doc)
    add_cover(doc)
    add_overview(doc)
    add_demo_and_future(doc)
    add_fill_in_answers(doc)

    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.text = ""
    r = header.add_run("IdeaPilot 作品介绍")
    set_run_font(r, size=9, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run("第 ")
    set_run_font(r, size=9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)
    r = footer.add_run(" 页")
    set_run_font(r, size=9, color=MUTED)

    doc.save(DOCX_PATH)
    MD_PATH.write_text(build_markdown(), encoding="utf-8")
    print(DOCX_PATH)
    print(MD_PATH)


if __name__ == "__main__":
    main()
